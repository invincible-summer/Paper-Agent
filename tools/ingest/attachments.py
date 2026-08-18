"""Unified uploaded-attachment storage and on-demand multimodal understanding.

Upload is deliberately cheap: persist the original, extract deterministic text,
and defer PDF layout/OCR and all VLM work until a deep-read or modality question.
Original bytes never enter history JSON; callers persist only returned metadata.
"""
from __future__ import annotations

import hashlib
import io
import logging
import mimetypes
import re
import shutil
import uuid
import zipfile
from pathlib import Path
from typing import Any

from core.blocking import run_cpu_bound
from core.config import get_settings
from core.models import Paper
from core.storage_context import StorageContext
from tools.pdf.structure.models import PaperElement, ParsedPaperDocument

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
UPLOAD_DIR = PROJECT_ROOT / "data" / "uploads"
MAX_FILE_BYTES = 50 * 1024 * 1024
TEXT_EXTENSIONS = {"txt", "md", "markdown", "tex", "bib"}
DOCUMENT_EXTENSIONS = {"pdf", "docx"}
API_DEFERRED_EXTENSIONS = {"doc", "xls", "xlsx"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "webp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS
_KNOWN_EXTENSIONS = SUPPORTED_EXTENSIONS | API_DEFERRED_EXTENSIONS
MULTIMODAL_EXTENSIONS = DOCUMENT_EXTENSIONS | IMAGE_EXTENSIONS
_ID_RE = re.compile(r"[a-f0-9]{32}")
_MARKER = "\n\n<!-- attachment-multimodal-understanding -->\n"

_MEDIA_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "doc": "application/msword",
    "xls": "application/vnd.ms-excel",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "tex": "application/x-tex",
    "txt": "text/plain",
    "md": "text/markdown",
    "markdown": "text/markdown",
    "bib": "application/x-bibtex",
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "webp": "image/webp",
}


class AttachmentError(ValueError):
    """User-readable attachment validation or extraction error."""


def normalize_ext(filename: str, content_type: str = "", raw: bytes = b"") -> str:
    ext = Path(filename or "").suffix.lower().lstrip(".")
    if ext == "markdown":
        return "md"
    if ext:
        # A caller-provided filename is authoritative. In particular, never
        # reinterpret an explicitly unsupported .ppt/.pptx (or arbitrary
        # suffix) as PDF/Word merely because a remote server reports a loose
        # Content-Type.
        return ext if ext in _KNOWN_EXTENSIONS else ""
    ct = (content_type or "").split(";", 1)[0].strip().lower()
    by_ct = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "application/msword": "doc",
        "application/vnd.ms-excel": "xls",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
        "image/png": "png", "image/jpeg": "jpg", "image/webp": "webp",
        "text/markdown": "md", "text/plain": "txt", "application/x-tex": "tex",
        "application/x-bibtex": "bib",
    }
    if ct in by_ct:
        return by_ct[ct]
    if raw.startswith(b"%PDF-"):
        return "pdf"
    if raw.startswith(b"\x89PNG\r\n\x1a\n"):
        return "png"
    if raw.startswith(b"\xff\xd8\xff"):
        return "jpg"
    if raw[:4] == b"RIFF" and raw[8:12] == b"WEBP":
        return "webp"
    if raw[:2] == b"PK":
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as zf:
                if "word/document.xml" in zf.namelist():
                    return "docx"
        except (OSError, zipfile.BadZipFile):
            pass
    return ""


def media_type_for(ext: str) -> str:
    return _MEDIA_TYPES.get(ext, mimetypes.guess_type(f"x.{ext}")[0] or "application/octet-stream")


def _effective_max_bytes(
    storage_context: StorageContext | None, max_bytes: int | None,
) -> int:
    if max_bytes is not None:
        return max_bytes
    if storage_context is not None and storage_context.channel == "openai_api":
        from core.api_storage_store import ApiStorageStore
        store = ApiStorageStore(storage_context)
        store.initialize()
        return store.get_policy().max_upload_bytes
    return MAX_FILE_BYTES


def original_path(attachment_id: str, ext: str) -> Path:
    return UPLOAD_DIR / f"{attachment_id}.{ext}"


def text_path(attachment_id: str) -> Path:
    return UPLOAD_DIR / f"{attachment_id}.txt"


def find_original(attachment_id: str) -> Path | None:
    if not _ID_RE.fullmatch(attachment_id or ""):
        return None
    # Every attachment also has ``<id>.txt`` as its extracted-text sidecar.
    # Probe real binary/document formats first so a WebP original is not hidden
    # by that sidecar; text formats come last (for a genuine TXT upload the
    # original and sidecar intentionally share the same path).
    priority = ["pdf", "docx", "doc", "xls", "xlsx", "png", "jpg", "jpeg", "webp", "tex", "md", "bib", "txt"]
    for ext in priority:
        fp = original_path(attachment_id, ext)
        if fp.is_file():
            return fp
    return None


def _extract_docx_text(raw: bytes) -> str:
    try:
        import docx
    except ImportError as e:  # pragma: no cover - dependency is required by project
        raise AttachmentError("docx 支持缺少 python-docx 依赖") from e
    try:
        document = docx.Document(io.BytesIO(raw))
    except Exception as e:  # noqa: BLE001
        raise AttachmentError("DOCX 文件损坏或不是有效的 Word 文档") from e
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table_no, table in enumerate(document.tables, 1):
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
        if not rows:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        blocks.append(f"[Table {table_no}]\n| " + " | ".join(rows[0]) + " |")
        blocks.append("| " + " | ".join(["---"] * width) + " |")
        blocks.extend("| " + " | ".join(r) + " |" for r in rows[1:])
    return "\n\n".join(blocks)


def extract_quick_text(raw: bytes, ext: str) -> str:
    """Fast deterministic extraction; scanned PDFs/images legitimately return ''."""
    if ext == "pdf":
        from tools.ingest.downloader import _extract_pdf
        try:
            return _extract_pdf(raw)
        except Exception:  # noqa: BLE001
            return ""
    if ext == "docx":
        return _extract_docx_text(raw)
    if ext in TEXT_EXTENSIONS:
        for enc in ("utf-8", "gb18030", "latin-1"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
    if ext in IMAGE_EXTENSIONS or ext in API_DEFERRED_EXTENSIONS:
        return ""
    raise AttachmentError(f"暂不支持的文件类型：{ext or '未知'}")


def _extract_quick_text_path(source: Path, ext: str) -> str:
    if ext == "pdf":
        from tools.pdf.structure.pymupdf_parser import PyMuPDFStructureParser
        return PyMuPDFStructureParser().parse(str(source)).raw_text
    if ext == "docx":
        try:
            import docx
            document = docx.Document(str(source))
        except Exception as exc:  # noqa: BLE001
            raise AttachmentError("DOCX 文件损坏或不是有效的 Word 文档") from exc
        blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_no, table in enumerate(document.tables, 1):
            rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells] for row in table.rows]
            if rows:
                width = max(len(row) for row in rows)
                rows = [row + [""] * (width - len(row)) for row in rows]
                blocks.append(f"[Table {table_no}]\n| " + " | ".join(rows[0]) + " |")
                blocks.append("| " + " | ".join(["---"] * width) + " |")
                blocks.extend("| " + " | ".join(row) + " |" for row in rows[1:])
        return "\n\n".join(blocks)
    if ext in TEXT_EXTENSIONS:
        raw = source.read_bytes()
        return extract_quick_text(raw, ext)
    if ext in IMAGE_EXTENSIONS or ext in API_DEFERRED_EXTENSIONS:
        return ""
    raise AttachmentError(f"暂不支持的文件类型：{ext or '未知'}")


def save_attachment_from_path(
    source: Path, filename: str, *, content_type: str = "", attachment_id: str = "",
    storage_context: StorageContext | None = None, session_id: str = "",
    owner_id: str = "", max_bytes: int | None = None,
) -> dict:
    max_bytes = _effective_max_bytes(storage_context, max_bytes)
    size = source.stat().st_size
    if size > max_bytes:
        from tools.ingest.downloader import format_byte_limit
        raise AttachmentError(f"文件超过 {format_byte_limit(max_bytes)} 上限")
    with source.open("rb") as stream:
        header = stream.read(4096)
    ext = normalize_ext(filename, content_type, header)
    is_api = storage_context is not None and storage_context.channel == "openai_api"
    allowed_extensions = SUPPORTED_EXTENSIONS | (API_DEFERRED_EXTENSIONS if is_api else set())
    if ext not in allowed_extensions:
        raise AttachmentError("暂不支持该文件类型；支持 PDF、DOCX、TEX、TXT、MD、BIB、PNG、JPG、WebP，API 另可暂存 DOC、XLS、XLSX")
    aid = attachment_id or uuid.uuid4().hex
    if not _ID_RE.fullmatch(aid):
        raise AttachmentError("附件 id 不合法")
    text = "" if ext in API_DEFERRED_EXTENSIONS else _extract_quick_text_path(source, ext)
    safe_name = Path(filename or f"upload.{ext}").name
    if not Path(safe_name).suffix:
        safe_name = f"{safe_name or 'upload'}.{ext}"
    extra: dict[str, str] = {}
    if storage_context is not None and storage_context.channel == "openai_api":
        if not session_id:
            raise AttachmentError("API 附件缺少会话隔离键")
        from core.api_artifact_store import ApiArtifactStore
        from core.api_storage_store import ApiStorageStore
        artifact_store = ApiArtifactStore(ApiStorageStore(storage_context))
        original = artifact_store.save_private_upload(
            source, session_id=session_id, logical_name=safe_name,
            mime_type=media_type_for(ext), category="upload",
        )
        sidecar = artifact_store.save_bytes(
            text.encode("utf-8"), category="upload_sidecar", scope=session_id,
            logical_name=f"{aid}.txt", mime_type="text/plain; charset=utf-8",
            ttl_seconds=artifact_store.storage.get_policy().upload_ttl_seconds,
        )
        now = __import__("time").time()
        with artifact_store.storage.connect() as conn:
            conn.execute(
                "INSERT OR IGNORE INTO api_session_artifacts(session_id, artifact_id, relation, created_at) VALUES (?, ?, 'upload_sidecar', ?)",
                (session_id, sidecar.id, now),
            )
            conn.commit()
        extra = {"artifact_id": original.id, "sidecar_artifact_id": sidecar.id,
                 "relative_path": original.relative_path, "text_relative_path": sidecar.relative_path}
    else:
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(source, original_path(aid, ext))
        text_path(aid).write_text(text, encoding="utf-8")
    status = ("deferred" if ext in API_DEFERRED_EXTENSIONS else
              "pending" if ext in MULTIMODAL_EXTENSIONS else "text_only")
    result = {
        "id": aid, "filename": safe_name, "char_count": len(text),
        "text_preview": text[:500], "ext": ext, "media_type": media_type_for(ext),
        "multimodal_status": status, "element_count": 0,
        "preview_url": f"/api/v1/chat/file/{aid}/raw" if ext in IMAGE_EXTENSIONS and storage_context is None else "",
        **extra,
    }
    if storage_context is None and owner_id:
        from core.web_artifact_store import register_web_artifact
        register_web_artifact("attachment", aid, owner_id, {
            key: result.get(key) for key in (
                "id", "filename", "char_count", "ext", "media_type",
                "multimodal_status", "element_count", "preview_url",
            )
        })
    return result


def save_attachment(
    raw: bytes, filename: str, *, content_type: str = "", attachment_id: str = "",
    storage_context: StorageContext | None = None, session_id: str = "",
    owner_id: str = "", max_bytes: int | None = None,
) -> dict:
    max_bytes = _effective_max_bytes(storage_context, max_bytes)
    if len(raw) > max_bytes:
        from tools.ingest.downloader import format_byte_limit
        raise AttachmentError(f"文件超过 {format_byte_limit(max_bytes)} 上限")
    ext = normalize_ext(filename, content_type, raw)
    is_api = storage_context is not None and storage_context.channel == "openai_api"
    allowed_extensions = SUPPORTED_EXTENSIONS | (API_DEFERRED_EXTENSIONS if is_api else set())
    if ext not in allowed_extensions:
        raise AttachmentError("暂不支持该文件类型；支持 PDF、DOCX、TEX、TXT、MD、BIB、PNG、JPG、WebP，API 另可暂存 DOC、XLS、XLSX")
    aid = attachment_id or uuid.uuid4().hex
    if not _ID_RE.fullmatch(aid):
        raise AttachmentError("附件 id 不合法")
    text = "" if ext in API_DEFERRED_EXTENSIONS else extract_quick_text(raw, ext)
    safe_name = Path(filename or f"upload.{ext}").name
    if not Path(safe_name).suffix:
        safe_name = f"{safe_name or 'upload'}.{ext}"
    extra: dict[str, str] = {}
    if storage_context is not None and storage_context.channel == "openai_api":
        if not session_id:
            raise AttachmentError("API 附件缺少会话隔离键")
        from core.api_artifact_store import ApiArtifactStore
        from core.api_storage_store import ApiStorageStore
        artifact_store = ApiArtifactStore(ApiStorageStore(storage_context))
        tmp = storage_context.temp_dir / f"upload-{uuid.uuid4().hex}.tmp"
        tmp.write_bytes(raw)
        try:
            original = artifact_store.save_private_upload(
                tmp, session_id=session_id, logical_name=safe_name,
                mime_type=media_type_for(ext), category="upload",
            )
        finally:
            tmp.unlink(missing_ok=True)
        sidecar = artifact_store.save_bytes(
            text.encode("utf-8"), category="upload_sidecar", scope=session_id,
            logical_name=f"{aid}.txt", mime_type="text/plain; charset=utf-8",
            ttl_seconds=artifact_store.storage.get_policy().upload_ttl_seconds,
        )
        now = __import__("time").time()
        with artifact_store.storage.connect() as conn:
            conn.execute(
                "INSERT OR IGNORE INTO api_session_artifacts(session_id, artifact_id, relation, created_at) VALUES (?, ?, 'upload_sidecar', ?)",
                (session_id, sidecar.id, now),
            )
            conn.commit()
        extra = {
            "artifact_id": original.id,
            "sidecar_artifact_id": sidecar.id,
            "relative_path": original.relative_path,
            "text_relative_path": sidecar.relative_path,
        }
    else:
        UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        original_path(aid, ext).write_bytes(raw)
        text_path(aid).write_text(text, encoding="utf-8")
    status = ("deferred" if ext in API_DEFERRED_EXTENSIONS else
              "pending" if ext in MULTIMODAL_EXTENSIONS else "text_only")
    result = {
        "id": aid,
        "filename": safe_name,
        "char_count": len(text),
        "text_preview": text[:500],
        "ext": ext,
        "media_type": media_type_for(ext),
        "multimodal_status": status,
        "element_count": 0,
        "preview_url": f"/api/v1/chat/file/{aid}/raw" if ext in IMAGE_EXTENSIONS and storage_context is None else "",
        **extra,
    }
    if storage_context is None and owner_id:
        from core.web_artifact_store import register_web_artifact
        register_web_artifact("attachment", aid, owner_id, {
            key: result.get(key) for key in (
                "id", "filename", "char_count", "ext", "media_type",
                "multimodal_status", "element_count", "preview_url",
            )
        })
    return result


def _api_artifact_path(attachment: dict, session, key: str) -> Path | None:
    artifact_id = attachment.get(key) or ""
    context = getattr(session, "storage_context", None)
    if not artifact_id or context is None or context.channel != "openai_api":
        return None
    from core.api_artifact_store import ApiArtifactStore
    from core.api_storage_store import ApiStorageStore
    return ApiArtifactStore(ApiStorageStore(context)).path_for(
        artifact_id, session_id=session.session_id
    )


def attachment_original_path(attachment: dict, session) -> Path | None:
    api_path = _api_artifact_path(attachment, session, "artifact_id")
    if api_path is not None:
        return api_path
    aid = attachment.get("id") or ""
    ext = (attachment.get("ext") or normalize_ext(attachment.get("filename") or "")).lower()
    return original_path(aid, ext) if ext else find_original(aid)


def attachment_text_path(attachment: dict, session) -> Path | None:
    api_path = _api_artifact_path(attachment, session, "sidecar_artifact_id")
    return api_path if api_path is not None else text_path(attachment.get("id") or "")


def attachment_document_id(attachment_id: str) -> str:
    return f"upload:{attachment_id}"


def attachment_by_id(session, attachment_id: str) -> dict | None:
    """Resolve a current-session attachment by stable id or exact filename.

    Tool schemas ask the model to pass the opaque id, but providers sometimes
    echo the user-visible filename instead. Accepting an exact, unambiguous
    filename keeps the tool robust without weakening session isolation: the
    result must still come from ``session.attachments`` and no filesystem path
    is consulted here.
    """
    ref = (attachment_id or "").strip()
    if not ref:
        return None
    attachments = list(session.attachments or [])
    by_id = next((a for a in attachments if a.get("id") == ref), None)
    if by_id is not None:
        return by_id
    matches = [a for a in attachments if (a.get("filename") or "").casefold() == ref.casefold()]
    return matches[0] if len(matches) == 1 else None


def session_element_scope(session) -> list[str]:
    return [p.id for p in session.all_papers()] + [
        attachment_document_id(a["id"]) for a in (session.attachments or []) if a.get("id")
    ]


def _understanding_text(elements: list[PaperElement]) -> str:
    parts: list[str] = []
    for el in elements:
        bits = [f"[{el.kind} {el.ordinal}]", el.caption or ""]
        dl = el.docling_extract or {}
        if dl.get("markdown"):
            bits.append(str(dl["markdown"]))
        if dl.get("latex"):
            bits.append(str(dl["latex"]))
        und = el.understanding or {}
        for key in ("description", "meaning", "role_in_paper", "role", "markdown", "latex"):
            if und.get(key):
                bits.append(str(und[key]))
        for key in ("components", "relations", "variables", "key_values", "trends"):
            if und.get(key):
                bits.append(str(und[key]))
        parts.append("\n".join(b for b in bits if b))
    return "\n\n".join(parts)


def _write_enriched_text(attachment: dict, session, base_text: str, elements: list[PaperElement]) -> str:
    base = (base_text or "").split(_MARKER, 1)[0].rstrip()
    semantic = _understanding_text(elements).strip()
    enriched = base + (_MARKER + semantic if semantic else "")
    context = getattr(session, "storage_context", None)
    if context is not None and context.channel == "openai_api":
        from core.api_artifact_store import ApiArtifactStore
        from core.api_storage_store import ApiStorageStore
        artifact_store = ApiArtifactStore(ApiStorageStore(context))
        sidecar = artifact_store.save_bytes(
            enriched.encode("utf-8"), category="upload_sidecar", scope=session.session_id,
            logical_name=f"{attachment.get('id', 'upload')}.txt",
            mime_type="text/plain; charset=utf-8",
            ttl_seconds=artifact_store.storage.get_policy().upload_ttl_seconds,
        )
        now = __import__("time").time()
        with artifact_store.storage.connect() as conn:
            conn.execute(
                "INSERT OR IGNORE INTO api_session_artifacts(session_id, artifact_id, relation, created_at) VALUES (?, ?, 'upload_sidecar', ?)",
                (session.session_id, sidecar.id, now),
            )
            conn.commit()
        attachment["sidecar_artifact_id"] = sidecar.id
        attachment["text_relative_path"] = sidecar.relative_path
    else:
        text_path(attachment.get("id") or "").write_text(enriched, encoding="utf-8")
    return enriched


def _docx_elements(raw: bytes, document_id: str, assets_dir: str) -> list[PaperElement]:
    root = Path(assets_dir)
    from tools.pdf.fetcher import _sanitize_filename_component as sanitize
    target = root / sanitize(document_id)
    target.mkdir(parents=True, exist_ok=True)
    elements: list[PaperElement] = []
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as zf:
            names = sorted(n for n in zf.namelist() if n.startswith("word/media/") and not n.endswith("/"))
            for name in names:
                try:
                    blob = zf.read(name)
                    ext = normalize_ext(Path(name).name, raw=blob)
                    if ext not in IMAGE_EXTENSIONS:
                        continue
                    ordinal = len(elements) + 1
                    out = target / f"figure_{ordinal}.{ext}"
                    out.write_bytes(blob)
                    elements.append(PaperElement(
                        element_id=f"{document_id}::figure::{ordinal}", kind="figure",
                        ordinal=ordinal, page=0, section="DOCX embedded media",
                        caption=f"Embedded image {ordinal}", asset_path=str(out),
                        image_hash=hashlib.sha256(blob).hexdigest(),
                    ))
                except Exception as e:  # noqa: BLE001
                    logger.debug("DOCX embedded image skipped (%s): %s", name, e)
    except Exception as e:  # noqa: BLE001
        logger.warning("DOCX media extraction failed: %s", e)
    return elements


def _image_element(path: Path, document_id: str, caption: str = "", assets_dir: str | None = None) -> PaperElement:
    raw = path.read_bytes()
    from tools.pdf.fetcher import _sanitize_filename_component as sanitize
    target = Path(assets_dir or get_settings().reader.assets_dir) / sanitize(document_id)
    target.mkdir(parents=True, exist_ok=True)
    out = target / f"figure_1{path.suffix.lower()}"
    out.write_bytes(raw)
    return PaperElement(
        element_id=f"{document_id}::figure::1", kind="figure", ordinal=1,
        page=0, section="Uploaded image", caption=caption or path.name,
        asset_path=str(out), image_hash=hashlib.sha256(raw).hexdigest(),
    )


def _prepare_docx_document(
    source: Path, document_id: str, assets_dir: str,
) -> ParsedPaperDocument:
    """Build DOCX text/elements synchronously for ``run_cpu_bound``."""
    raw = source.read_bytes()
    return ParsedPaperDocument(
        raw_text=_extract_docx_text(raw),
        elements=_docx_elements(raw, document_id, assets_dir),
        doc_fingerprint=hashlib.sha256(raw).hexdigest(),
        parser_backend="docx",
    )


def _prepare_image_document(
    source: Path, document_id: str, title: str, assets_dir: str,
) -> ParsedPaperDocument:
    """Build a standalone-image document synchronously for ``run_cpu_bound``."""
    element = _image_element(source, document_id, title, assets_dir)
    raw = source.read_bytes()
    return ParsedPaperDocument(
        elements=[element], doc_fingerprint=hashlib.sha256(raw).hexdigest(),
        parser_backend="image", page_count=1,
    )


def _secure_element_assets(context, elements: list[PaperElement]) -> None:
    for element in elements:
        if element.asset_path and Path(element.asset_path).is_file():
            try:
                context.secure_private_file(element.asset_path)
            except Exception:  # noqa: BLE001
                pass


async def ensure_attachment_understood(attachment: dict, session, *, focus: str = "") -> dict:
    """Run cached multimodal understanding for one current-session attachment."""
    aid = attachment.get("id") or ""
    ext = (attachment.get("ext") or "").lower()
    if not ext:
        # Old history rows did not persist ``ext``. Recover it from the original
        # filename before looking at disk, because ``<id>.txt`` may only be the
        # extracted-text sidecar of a legacy PDF/DOCX/image upload.
        ext = normalize_ext(attachment.get("filename") or "")
    source = attachment_original_path(attachment, session)
    if source is not None and not source.is_file():
        source = None
    if not ext and source:
        ext = source.suffix.lower().lstrip(".")
    if ext:
        attachment["ext"] = ext
        attachment["media_type"] = media_type_for(ext)
    if ext in API_DEFERRED_EXTENSIONS:
        attachment["multimodal_status"] = "deferred"
        return {
            "id": aid,
            "status": "deferred",
            "element_count": 0,
            "reason": "该格式已安全保存，但当前版本尚未提供解析能力。",
        }
    if ext not in MULTIMODAL_EXTENSIONS:
        attachment["multimodal_status"] = "text_only" if ext in TEXT_EXTENSIONS or not source else "unsupported"
        return {"id": aid, "status": attachment["multimodal_status"], "element_count": 0}
    if source is None:
        # Legacy records only retained extracted text; keep them queryable.
        attachment["multimodal_status"] = "legacy_text_only"
        return {"id": aid, "status": "legacy_text_only", "element_count": 0}
    if attachment.get("multimodal_status") == "ready":
        return {"id": aid, "status": "ready", "element_count": int(attachment.get("element_count") or 0)}

    document_id = attachment_document_id(aid)
    title = attachment.get("filename") or source.name
    paper = Paper(id=document_id, title=title, source="upload", pdf_path=str(source) if ext == "pdf" else None)
    doc: ParsedPaperDocument | None = None
    elements: list[PaperElement] = []
    recovered_text = ""
    try:
        from agents.reader_agent import (
            _understand_and_persist_elements, parse_and_understand,
            persist_api_element_assets,
        )
        from tools.storage.database import Database
        context = getattr(session, "storage_context", None)
        db = Database(storage_context=context) if context is not None else Database(get_settings().storage.sqlite_path)
        assets_dir = str(context.blob_dir / "element_assets") if context is not None and context.channel == "openai_api" else get_settings().reader.assets_dir
        try:
            if ext == "pdf":
                doc = await (
                    parse_and_understand(
                        paper, db, assets_dir=assets_dir, storage_context=context,
                        session_id=session.session_id,
                    )
                    if context is not None else parse_and_understand(paper, db, assets_dir=assets_dir)
                )
                if doc:
                    elements = doc.elements
                    recovered_text = doc.raw_text or ""
            elif ext == "docx":
                doc = await run_cpu_bound(
                    _prepare_docx_document, source, document_id, assets_dir
                )
                recovered_text = doc.raw_text or ""
                elements = doc.elements
                if context is not None and context.channel == "openai_api":
                    await run_cpu_bound(_secure_element_assets, context, elements)
                    await run_cpu_bound(
                        persist_api_element_assets, paper, doc, context,
                        session_id=session.session_id,
                    )
                if elements:
                    await (_understand_and_persist_elements(paper, doc, db, storage_context=context) if context is not None else _understand_and_persist_elements(paper, doc, db))
            else:
                doc = await run_cpu_bound(
                    _prepare_image_document, source, document_id, title, assets_dir
                )
                elements = doc.elements
                if context is not None and context.channel == "openai_api":
                    await run_cpu_bound(_secure_element_assets, context, elements)
                    await run_cpu_bound(
                        persist_api_element_assets, paper, doc, context,
                        session_id=session.session_id,
                    )
                await (_understand_and_persist_elements(paper, doc, db, storage_context=context) if context is not None else _understand_and_persist_elements(paper, doc, db))
        finally:
            db.close()

        sidecar_path = attachment_text_path(attachment, session)
        old_text = await run_cpu_bound(
            lambda: sidecar_path.read_text(encoding="utf-8")
            if sidecar_path and sidecar_path.is_file() else ""
        )
        enriched = await run_cpu_bound(
            _write_enriched_text, attachment, session, recovered_text or old_text, elements
        )
        from tools.storage.vectorstore import VectorStore

        def _reindex_attachment() -> None:
            try:
                VectorStore(
                    storage_context=getattr(session, "storage_context", None)
                ).upsert_text_chunks(
                    aid, title, enriched, session_id=session.session_id
                )
            except Exception as e:  # noqa: BLE001
                logger.debug("attachment chunk re-index skipped for %s: %s", aid, e)

        await run_cpu_bound(_reindex_attachment)
        attachment.update({
            "char_count": len(enriched), "text_preview": enriched[:500],
            "multimodal_status": "ready", "element_count": len(elements),
        })
        return {"id": aid, "status": "ready", "element_count": len(elements),
                "char_count": len(enriched), "elements": [e.short_ref() for e in elements]}
    except Exception as e:  # noqa: BLE001
        logger.warning("attachment understanding failed for %s: %s", aid, e)
        attachment["multimodal_status"] = "degraded"
        return {"id": aid, "status": "degraded", "element_count": len(elements), "error": str(e)}
