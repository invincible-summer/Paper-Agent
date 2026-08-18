"""Safe URL download + text extraction for multimodal file inputs.

清小搭 sends user files as URLs (never base64). This module downloads them
with the same SSRF guard as the PDF fetcher, caps size, and extracts text:
  pdf  -> PyMuPDF (column-aware parser)
  docx -> python-docx
  txt/md/markdown -> multi-encoding decode
"""
from __future__ import annotations

import logging
import os
import tempfile
import uuid
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx

from tools.pdf.fetcher import _is_safe_url

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 50 * 1024 * 1024  # non-API compatibility default


class IngestError(ValueError):
    """Raised when a remote file cannot be downloaded or extracted."""


def format_byte_limit(max_bytes: int) -> str:
    """Human-readable configured limit without hard-coded 50 MB wording."""
    mib = max_bytes / (1024 * 1024)
    return f"{int(mib)} MiB" if mib.is_integer() else f"{mib:.1f} MiB"


def _ext_of(url: str, filename: str, content_type: str) -> str:
    for cand in (filename, Path(urlparse(url).path).name):
        ext = Path(cand or "").suffix.lower().lstrip(".")
        if ext:
            return ext
    ct = (content_type or "").lower()
    if "pdf" in ct:
        return "pdf"
    if "wordprocessingml" in ct or "msword" in ct:
        return "docx"
    if "markdown" in ct:
        return "md"
    if "text/plain" in ct:
        return "txt"
    return ""


async def download_to_temp(
    url: str, *, temp_dir: str | Path | None = None,
    max_bytes: int = MAX_FILE_BYTES, max_redirects: int = 5,
) -> tuple[Path, str]:
    """Stream a public URL to a private temp file with per-hop SSRF checks."""
    directory = Path(temp_dir) if temp_dir is not None else Path(tempfile.gettempdir())
    directory.mkdir(parents=True, exist_ok=True)
    target = directory / f"ingest-{uuid.uuid4().hex}.tmp"
    current = url
    try:
        async with httpx.AsyncClient(timeout=60, follow_redirects=False,
                                     proxy=None, trust_env=False) as client:
            for hop in range(max_redirects + 1):
                safe, reason = _is_safe_url(current)
                if not safe:
                    raise IngestError(f"URL 不安全，已拒绝下载：{reason}")
                async with client.stream("GET", current) as resp:
                    if resp.status_code in {301, 302, 303, 307, 308}:
                        location = resp.headers.get("location")
                        if not location or hop >= max_redirects:
                            raise IngestError("下载重定向过多或缺少 Location")
                        current = urljoin(current, location)
                        continue
                    if resp.status_code != 200:
                        raise IngestError(f"下载失败（HTTP {resp.status_code}）")
                    content_type = resp.headers.get("content-type", "")
                    declared = resp.headers.get("content-length")
                    if declared and declared.isdigit() and int(declared) > max_bytes:
                        raise IngestError(f"文件超过 {format_byte_limit(max_bytes)} 上限")
                    total = 0
                    with target.open("xb") as output:
                        async for chunk in resp.aiter_bytes():
                            if not chunk:
                                continue
                            total += len(chunk)
                            if total > max_bytes:
                                raise IngestError(f"文件超过 {format_byte_limit(max_bytes)} 上限")
                            output.write(chunk)
                        output.flush()
                        os.fsync(output.fileno())
                    os.chmod(target, 0o600)
                    return target, content_type
        raise IngestError("下载重定向失败")
    except IngestError:
        target.unlink(missing_ok=True)
        raise
    except Exception as exc:  # noqa: BLE001
        target.unlink(missing_ok=True)
        raise IngestError("下载出错") from exc


async def download_bytes(url: str, max_bytes: int = MAX_FILE_BYTES) -> tuple[bytes, str]:
    """Compatibility helper; API ingestion uses ``download_to_temp`` directly."""
    path, content_type = await download_to_temp(url, max_bytes=max_bytes)
    try:
        return path.read_bytes(), content_type
    finally:
        path.unlink(missing_ok=True)


def extract_text(raw: bytes, ext: str) -> str:
    """Extract text from downloaded bytes. Raises IngestError when unsupported."""
    ext = (ext or "").lower()
    if ext == "pdf" or raw[:5] == b"%PDF-":
        return _extract_pdf(raw)
    if ext == "docx":
        return _extract_docx(raw)
    if ext in ("png", "jpg", "jpeg", "webp"):
        return ""
    if ext in ("txt", "md", "markdown", "tex", "bib", ""):
        for enc in ("utf-8", "gb18030", "latin-1"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
    if ext == "doc":
        # Legacy Word binary — no libreoffice/antiword/pandoc in this env, and
        # python-docx only reads OOXML (.docx). Honest degrade with an action.
        raise IngestError("旧版 .doc（Word 二进制）当前环境无法解析，请在 Word 里另存为 .docx 后上传")
    raise IngestError(f"暂不支持的文件类型：{ext or '未知'}（支持 pdf/docx/tex/txt/md/bib）")


def _extract_pdf(raw: bytes) -> str:
    from tools.pdf.structure.pymupdf_parser import PyMuPDFStructureParser

    tmp = None
    try:
        with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
            f.write(raw)
            tmp = f.name
        # No paper_id => no element assets written; we only need raw text for a
        # chat file upload. The PyMuPDF structure backend (not Docling) keeps
        # this synchronous and fast, with column-aware text + header/footer
        # stripping — the same quality the old parser had, plus page numbers.
        return PyMuPDFStructureParser().parse(tmp).raw_text
    finally:
        if tmp:
            Path(tmp).unlink(missing_ok=True)


def _extract_docx(raw: bytes) -> str:
    import io

    try:
        import docx
    except ImportError as e:
        raise IngestError("docx 支持缺少 python-docx 依赖") from e
    document = docx.Document(io.BytesIO(raw))
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table_no, table in enumerate(document.tables, 1):
        rows = [[cell.text.strip().replace("\n", " ") for cell in row.cells]
                for row in table.rows]
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        blocks.append(f"[Table {table_no}]\n| " + " | ".join(rows[0]) + " |")
        blocks.append("| " + " | ".join(["---"] * width) + " |")
        blocks.extend("| " + " | ".join(row) + " |" for row in rows[1:])
    return "\n\n".join(blocks)


async def download_and_extract(url: str, filename: str = "") -> tuple[str, str]:
    """Full pipeline: download -> detect type -> extract text.

    Returns (text, ext). Raises IngestError with a user-readable reason.
    """
    raw, content_type = await download_bytes(url)
    ext = _ext_of(url, filename, content_type)
    return extract_text(raw, ext), ext
