"""manuscript_export: md/docx/tex manuscript files + docx/tex upload parsing."""
import docx
import zipfile

from tools.writing.manuscript_export import export_manuscript

CONTENT = """## 引言
这是**引言**段落，含加粗。

- 要点一
- 要点二

| 方法 | 指标 |
|---|---|
| A | 0.9 |

## 方法
方法内容。
"""


def test_export_md(tmp_path, monkeypatch):
    import tools.writing.manuscript_export as me
    monkeypatch.setattr(me, "_EXPORT_DIR", tmp_path)
    rec = export_manuscript("测试稿", CONTENT, "md")
    text = (tmp_path / rec["fileName"]).read_text(encoding="utf-8")
    assert text.startswith("# 测试稿")
    assert rec["fileType"] == "text"


def test_export_docx_roundtrip(tmp_path, monkeypatch):
    import tools.writing.manuscript_export as me
    monkeypatch.setattr(me, "_EXPORT_DIR", tmp_path)
    rec = export_manuscript("测试稿", CONTENT, "docx")
    doc = docx.Document(str(tmp_path / rec["fileName"]))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert "引言" in heads and "方法" in heads
    assert rec["mimeType"].endswith("wordprocessingml.document")


def test_export_docx_preserves_markdown_links_as_word_hyperlinks(tmp_path, monkeypatch):
    import tools.writing.manuscript_export as me
    monkeypatch.setattr(me, "_EXPORT_DIR", tmp_path)
    rec = export_manuscript(
        "链接稿", "引用 [论文落地页](https://example.org/paper) 与 **重点**。", "docx"
    )
    with zipfile.ZipFile(tmp_path / rec["fileName"]) as archive:
        document = archive.read("word/document.xml").decode("utf-8")
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
    assert "论文落地页" in document
    assert "https://example.org/paper" in rels
    assert "w:hyperlink" in document
    assert "重点" in document


def test_export_tex_structure(tmp_path, monkeypatch):
    import tools.writing.manuscript_export as me
    monkeypatch.setattr(me, "_EXPORT_DIR", tmp_path)
    rec = export_manuscript("测试稿", CONTENT, "tex")
    text = (tmp_path / rec["fileName"]).read_text(encoding="utf-8")
    assert "\\subsection{引言}" in text and "\\subsection{方法}" in text
    assert "\\begin{itemize}" in text and "\\item 要点一" in text
    assert "\\textbf{引言}段落" in text  # inline bold preserved
    assert "ctexart" in text


def test_export_invalid_format_falls_back_md(tmp_path, monkeypatch):
    import tools.writing.manuscript_export as me
    monkeypatch.setattr(me, "_EXPORT_DIR", tmp_path)
    rec = export_manuscript("x", CONTENT, "pdf")
    assert rec["fileName"].endswith(".md")


def test_docx_upload_parsing():
    """tools/ingest extracts docx text (chat upload path delegates here)."""
    import io
    d = docx.Document()
    d.add_heading("我的论文", level=1)
    d.add_paragraph("摘要内容")
    buf = io.BytesIO()
    d.save(buf)
    from tools.ingest.downloader import extract_text
    text = extract_text(buf.getvalue(), "docx")
    assert "我的论文" in text and "摘要内容" in text
    # tex decodes as plain text
    assert "\\section{引言}" in extract_text("\\section{引言}".encode(), "tex")


def test_export_docx_preserves_ordered_lists_and_escaped_table_pipes(tmp_path, monkeypatch):
    import tools.writing.manuscript_export as me
    monkeypatch.setattr(me, "_EXPORT_DIR", tmp_path)
    content = r"""# 结构

1. 第一步
2. 第二步

- 无序项

| 名称 | 说明 |
|---|---|
| A | 包含 \| 管道与 **重点** |
"""
    rec = export_manuscript("超长中文综述标题用于验证安全截断与Unicode导出能力" * 3, content, "docx")
    doc = docx.Document(str(tmp_path / rec["fileName"]))
    numbered = [p.text for p in doc.paragraphs if p.style.name == "List Number"]
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert numbered == ["第一步", "第二步"]
    assert bullets == ["无序项"]
    assert any("包含 | 管道与 重点" in cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert (tmp_path / rec["fileName"]).read_bytes().startswith(b"PK")
