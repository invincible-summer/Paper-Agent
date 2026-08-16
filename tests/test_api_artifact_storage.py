"""Module 3: API artifact, upload, metadata/Chroma and export isolation."""
from __future__ import annotations

import asyncio
import sqlite3
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from agents.session import ChatSession
from core.api_artifact_store import ApiArtifactStore
from core.api_storage_store import ApiStorageStore
from core.storage_context import StorageContext


def _artifact_store(tmp_path: Path) -> ApiArtifactStore:
    context = StorageContext.openai_api(root_dir=tmp_path / "openai-api")
    return ApiArtifactStore(ApiStorageStore(context))



def _ensure_session(store: ApiArtifactStore, session_id: str) -> None:
    import time
    now = time.time()
    with store.storage.connect() as conn:
        conn.execute(
            "INSERT OR IGNORE INTO api_sessions"
            "(id, credential_id, rag_session_id, created_at, updated_at, last_accessed_at, expires_at)"
            " VALUES (?, 'test-key', ?, ?, ?, ?, ?)",
            (session_id, session_id, now, now, now, now + 86400),
        )
        conn.commit()

def test_public_pdf_deduplicates_but_private_uploads_do_not_cross_sessions(tmp_path: Path):
    store = _artifact_store(tmp_path)
    source = tmp_path / "paper.pdf"
    source.write_bytes(b"%PDF-1.4\nsame")
    public_a = store.save_public_pdf(source, logical_name="paper-a.pdf")
    public_b = store.save_public_pdf(source, logical_name="paper-b.pdf")
    assert public_a.id == public_b.id
    assert public_a.path == public_b.path

    _ensure_session(store, "session-a")
    _ensure_session(store, "session-b")
    private_a = store.save_private_upload(
        source, session_id="session-a", logical_name="private.pdf",
        mime_type="application/pdf",
    )
    private_b = store.save_private_upload(
        source, session_id="session-b", logical_name="private.pdf",
        mime_type="application/pdf",
    )
    assert private_a.id != private_b.id
    assert private_a.path != private_b.path
    assert store.path_for(private_a.id, session_id="session-a") == private_a.path
    assert store.path_for(private_a.id, session_id="session-b") is None


def test_exports_have_unique_public_aliases_and_expire_independently(tmp_path: Path):
    store = _artifact_store(tmp_path)
    source = tmp_path / "same.md"
    source.write_text("# same", encoding="utf-8")
    _ensure_session(store, "session-a")
    first = store.save_export(
        source, session_id="session-a", display_name="中文报告.md", mime_type="text/markdown"
    )
    second = store.save_export(
        source, session_id="session-a", display_name="中文报告.md", mime_type="text/markdown"
    )
    assert first.id != second.id
    assert first.public_alias != second.public_alias
    assert store.resolve_public_alias(first.public_alias or "") is not None
    with store.storage.connect() as conn:
        conn.execute("UPDATE api_artifacts SET expires_at = 0 WHERE id = ?", (first.id,))
        conn.commit()
    assert store.resolve_public_alias(first.public_alias or "") is None
    assert store.resolve_public_alias(second.public_alias or "") is not None


def test_api_attachment_original_and_sidecar_live_only_under_api_root(tmp_path: Path):
    from tools.ingest.attachments import (
        attachment_original_path, attachment_text_path, save_attachment,
    )

    context = StorageContext.openai_api(root_dir=tmp_path / "openai-api")
    session = ChatSession(
        channel="openai_api", storage_context=context, session_id="session-upload"
    )
    _ensure_session(ApiArtifactStore(ApiStorageStore(context)), session.session_id)
    web_uploads = tmp_path / "web-uploads"
    import tools.ingest.attachments as attachments
    old_upload_dir = attachments.UPLOAD_DIR
    attachments.UPLOAD_DIR = web_uploads
    try:
        record = save_attachment(
            b"hello API attachment", "notes.txt",
            storage_context=context, session_id=session.session_id,
        )
    finally:
        attachments.UPLOAD_DIR = old_upload_dir
    session.attachments = [record]
    original = attachment_original_path(record, session)
    sidecar = attachment_text_path(record, session)
    assert original and original.is_file() and original.is_relative_to(context.root_dir)
    assert sidecar and sidecar.is_file() and sidecar.is_relative_to(context.root_dir)
    assert original.read_bytes() == b"hello API attachment"
    assert sidecar.read_text(encoding="utf-8") == "hello API attachment"
    assert not web_uploads.exists()
    assert record["relative_path"].startswith("blobs/upload/")
    assert record["text_relative_path"].startswith("blobs/upload_sidecar/")


def test_api_metadata_vision_cache_and_chroma_paths_are_isolated(tmp_path: Path):
    from core.multimodal.cache import VisionCache
    from tools.storage.database import Database
    from tools.storage.vectorstore import VectorStore

    context = StorageContext.openai_api(root_dir=tmp_path / "openai-api")
    web_db = tmp_path / "web" / "metadata.db"
    api_db = Database(storage_context=context)
    api_db.save_vision_cache("hash", "figure", 1, '{"description":"api"}')
    api_db.close()
    assert context.metadata_db.is_file()
    assert not web_db.exists()
    with sqlite3.connect(context.metadata_db) as conn:
        assert conn.execute("SELECT COUNT(*) FROM vision_cache").fetchone()[0] == 1
    assert VectorStore(storage_context=context)._path == str(context.chroma_dir)
    assert VectorStore(path=str(tmp_path / "web" / "chroma"))._path != str(context.chroma_dir)


def test_api_unicode_docx_export_downloads_with_pk_signature(tmp_path: Path, monkeypatch):
    from app.api.v1.files import get_file
    from tools.writing.manuscript_export import export_manuscript

    root = tmp_path / "openai-api"
    monkeypatch.setenv("OPENAI_API_STORAGE_ROOT", str(root))
    context = StorageContext.openai_api(root_dir=root)
    _ensure_session(ApiArtifactStore(ApiStorageStore(context)), "session-docx")
    record = export_manuscript(
        "堂吉诃德的理性与非理性——韦伯卡里斯马理论视角下的论文骨架",
        "# 绪论\n\n" + "这是用于验证 Word 下载的完整中文段落。" * 10,
        "docx", storage_context=context, session_id="session-docx",
    )
    response = get_file(record["fileName"])
    assert response.status_code == 200
    assert response.media_type.startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert Path(response.path).read_bytes().startswith(b"PK")
    disposition = response.headers["content-disposition"].lower()
    assert "filename*=utf-8''" in disposition


def test_web_export_path_remains_unchanged(tmp_path: Path, monkeypatch):
    from tools.writing import manuscript_export

    monkeypatch.setattr(manuscript_export, "_EXPORT_DIR", tmp_path / "web-exports")
    record = manuscript_export.export_manuscript(
        "Web 文稿", "# 正文\n\n" + "web data " * 20, "md"
    )
    assert (tmp_path / "web-exports" / record["fileName"]).is_file()
    assert "displayName" not in record


@pytest.mark.anyio
async def test_openai_media_registration_uses_api_private_artifacts_without_vlm(tmp_path: Path, monkeypatch):
    import base64
    from app.api.v1.multimodal import MediaPart, process_media_parts

    context = StorageContext.openai_api(root_dir=tmp_path / "openai-api")
    artifact_store = ApiArtifactStore(ApiStorageStore(context))
    _ensure_session(artifact_store, "session-image")
    png = base64.b64encode(b"\x89PNG\r\n\x1a\nimage-bytes").decode("ascii")
    notes, attachments, errors = await process_media_parts(
        [MediaPart(type="image", url=f"data:image/png;base64,{png}", filename="图像.png")],
        storage_context=context,
        session_id="session-image",
    )
    assert errors == [] and notes
    record = attachments[0]
    assert record["multimodal_status"] == "pending"
    assert record["artifact_id"] and record["sidecar_artifact_id"]
    assert artifact_store.path_for(record["artifact_id"], session_id="session-image").read_bytes().startswith(b"\x89PNG")
    # Registration is intentionally zero-VLM: it has no understanding fields or element rows.
    with artifact_store.storage.connect() as conn:
        assert conn.execute("SELECT COUNT(*) FROM api_artifacts WHERE category='element_asset'").fetchone()[0] == 0


def test_api_private_attachment_checkpoint_restores_artifact_references(tmp_path: Path):
    from core.api_checkpoint import ApiCheckpointStore, ApiCredentialPrincipal
    from tools.ingest.attachments import save_attachment

    context = StorageContext.openai_api(root_dir=tmp_path / "openai-api")
    checkpoint = ApiCheckpointStore(ApiStorageStore(context))
    principal = ApiCredentialPrincipal("key", "key", "admin", "database")
    loaded = checkpoint.get_or_create(principal, "user", [{"role": "user", "content": "upload"}])
    record = save_attachment(
        b"checkpoint attachment", "notes.txt", storage_context=context,
        session_id=loaded.session_id,
    )
    loaded.session.attachments = [record]
    checkpoint.save_checkpoint_sync(loaded.session)
    restored = ApiCheckpointStore(ApiStorageStore(context))._load_checkpoint_sync(loaded.session_id)
    assert restored is not None
    assert restored.attachments[0]["artifact_id"] == record["artifact_id"]
    assert restored.attachments[0]["relative_path"] == record["relative_path"]


def test_expired_api_alias_is_404_even_if_web_file_has_same_name(tmp_path: Path, monkeypatch):
    from app.api.v1 import files as files_api
    from fastapi import HTTPException

    root = tmp_path / "openai-api"
    web_exports = tmp_path / "web-exports"
    web_exports.mkdir()
    monkeypatch.setenv("OPENAI_API_STORAGE_ROOT", str(root))
    monkeypatch.setattr(files_api, "_EXPORT_DIR", web_exports)
    context = StorageContext.openai_api(root_dir=root)
    store = ApiArtifactStore(ApiStorageStore(context))
    _ensure_session(store, "session-expired")
    source = tmp_path / "x.md"
    source.write_text("api", encoding="utf-8")
    artifact = store.save_export(
        source, session_id="session-expired", display_name="x.md", mime_type="text/markdown"
    )
    (web_exports / artifact.public_alias).write_text("web collision", encoding="utf-8")
    with store.storage.connect() as conn:
        conn.execute("UPDATE api_artifacts SET expires_at = 0 WHERE id = ?", (artifact.id,))
        conn.commit()
    with pytest.raises(HTTPException) as exc:
        files_api.get_file(artifact.public_alias)
    assert exc.value.status_code == 404


@pytest.mark.anyio
async def test_api_png_lazy_multimodal_uses_api_assets_metadata_and_chroma(monkeypatch, tmp_path: Path):
    import agents.reader_agent as reader
    import tools.ingest.attachments as attachments
    import tools.storage.vectorstore as vectorstore

    context = StorageContext.openai_api(root_dir=tmp_path / "openai-api")
    checkpoint_store = __import__("core.api_checkpoint", fromlist=["ApiCheckpointStore"]).ApiCheckpointStore(
        ApiStorageStore(context)
    )
    principal_cls = __import__("core.api_checkpoint", fromlist=["ApiCredentialPrincipal"]).ApiCredentialPrincipal
    loaded = checkpoint_store.get_or_create(
        principal_cls("key", "key", "admin", "database"), None,
        [{"role": "user", "content": "image"}],
    )
    record = attachments.save_attachment(
        b"\x89PNG\r\n\x1a\nimage", "architecture.png",
        storage_context=context, session_id=loaded.session_id,
    )
    loaded.session.attachments = [record]
    calls = []

    async def fake_persist(paper, doc, db, *, storage_context=None):
        calls.append(storage_context)
        doc.elements[0].understanding = {"description": "architecture"}
        db.save_elements(paper.id, doc.elements, doc.doc_fingerprint)

    class FakeVectorStore:
        def __init__(self, *args, **kwargs):
            calls.append(kwargs.get("storage_context"))
        def upsert_text_chunks(self, *args, **kwargs):
            return 1

    monkeypatch.setattr(reader, "_understand_and_persist_elements", fake_persist)
    monkeypatch.setattr(vectorstore, "VectorStore", FakeVectorStore)
    result = await attachments.ensure_attachment_understood(record, loaded.session)
    assert result["status"] == "ready" and result["element_count"] == 1
    assert context in calls
    with sqlite3.connect(context.metadata_db) as conn:
        row = conn.execute("SELECT asset_path, understanding FROM paper_elements").fetchone()
    assert row and Path(row[0]).is_relative_to(context.root_dir)
    assert Path(row[0]).stat().st_mode & 0o777 == 0o600
    assert "architecture" in row[1]


@pytest.mark.anyio
async def test_api_docx_embedded_image_is_lazy_and_private(monkeypatch, tmp_path: Path):
    import zipfile
    import docx
    import agents.reader_agent as reader
    import tools.ingest.attachments as attachments
    import tools.storage.vectorstore as vectorstore

    source = tmp_path / "source.docx"
    document = docx.Document()
    document.add_paragraph("Document body for lazy multimodal extraction.")
    document.save(source)
    with zipfile.ZipFile(source, "a") as archive:
        archive.writestr("word/media/image99.png", b"\x89PNG\r\n\x1a\nembedded")

    context = StorageContext.openai_api(root_dir=tmp_path / "openai-api")
    from core.api_checkpoint import ApiCheckpointStore, ApiCredentialPrincipal
    checkpoint = ApiCheckpointStore(ApiStorageStore(context))
    loaded = checkpoint.get_or_create(
        ApiCredentialPrincipal("key", "key", "admin", "database"), None,
        [{"role": "user", "content": "docx"}],
    )
    record = attachments.save_attachment_from_path(
        source, "paper.docx", storage_context=context, session_id=loaded.session_id
    )
    loaded.session.attachments = [record]

    async def fake_persist(paper, parsed, db, *, storage_context=None):
        for element in parsed.elements:
            element.understanding = {"description": "embedded figure"}
        db.save_elements(paper.id, parsed.elements, parsed.doc_fingerprint)

    class FakeVectorStore:
        def __init__(self, *args, **kwargs): pass
        def upsert_text_chunks(self, *args, **kwargs): return 1

    monkeypatch.setattr(reader, "_understand_and_persist_elements", fake_persist)
    monkeypatch.setattr(vectorstore, "VectorStore", FakeVectorStore)
    result = await attachments.ensure_attachment_understood(record, loaded.session)
    assert result["status"] == "ready" and result["element_count"] >= 1
    assert all(Path(item["asset_path"]).is_relative_to(context.root_dir)
               for item in __import__("tools.storage.database", fromlist=["Database"]).Database(storage_context=context).get_elements(
                   f"upload:{record['id']}"
               ))


@pytest.mark.anyio
async def test_api_pdf_lazy_read_uses_shared_chokepoint_and_context(monkeypatch, tmp_path: Path):
    import fitz
    import agents.reader_agent as reader
    import tools.ingest.attachments as attachments
    import tools.storage.vectorstore as vectorstore
    from tools.pdf.structure.models import ParsedPaperDocument

    source = tmp_path / "paper.pdf"
    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "A valid PDF body for API lazy deep reading.")
    pdf.save(source)
    pdf.close()

    context = StorageContext.openai_api(root_dir=tmp_path / "openai-api")
    from core.api_checkpoint import ApiCheckpointStore, ApiCredentialPrincipal
    checkpoint = ApiCheckpointStore(ApiStorageStore(context))
    loaded = checkpoint.get_or_create(
        ApiCredentialPrincipal("key", "key", "admin", "database"), None,
        [{"role": "user", "content": "pdf"}],
    )
    record = attachments.save_attachment_from_path(
        source, "paper.pdf", storage_context=context, session_id=loaded.session_id
    )
    loaded.session.attachments = [record]
    seen = []

    async def fake_parse(paper, db, *, assets_dir=None, storage_context=None, session_id=""):
        seen.append((paper.pdf_path, assets_dir, storage_context))
        return ParsedPaperDocument(raw_text="Recovered API PDF text", doc_fingerprint="fp")

    class FakeVectorStore:
        def __init__(self, *args, **kwargs): pass
        def upsert_text_chunks(self, *args, **kwargs): return 1

    monkeypatch.setattr(reader, "parse_and_understand", fake_parse)
    monkeypatch.setattr(vectorstore, "VectorStore", FakeVectorStore)
    result = await attachments.ensure_attachment_understood(record, loaded.session)
    assert result["status"] == "ready"
    assert seen and Path(seen[0][0]).is_relative_to(context.root_dir)
    assert Path(seen[0][1]).is_relative_to(context.root_dir)
    assert seen[0][2] == context
