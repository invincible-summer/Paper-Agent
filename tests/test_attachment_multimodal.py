"""Uploaded-attachment multimodal pipeline tests (offline).

Upload itself is deterministic and VLM-free. PDF/image/DOCX visual work is
performed only when a current-session tool explicitly needs it, then cached in
session metadata and the global element store.
"""
from __future__ import annotations

import asyncio
import base64
import io
import zipfile
from pathlib import Path

import docx

from agents.session import ChatSession
from core.config import get_settings
from tools.pdf.structure.models import PaperElement, ParsedPaperDocument

# Minimal PNG bytes are sufficient here: ingestion validates the signature and
# the vision layer is stubbed before any decoder/provider sees the payload.
PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _patch_paths(monkeypatch, tmp_path):
    import tools.ingest.attachments as att

    uploads = tmp_path / "uploads"
    assets = tmp_path / "assets"
    db_path = str(tmp_path / "metadata.db")
    monkeypatch.setattr(att, "UPLOAD_DIR", uploads)
    monkeypatch.setattr(get_settings().reader, "assets_dir", str(assets))
    monkeypatch.setattr(get_settings().storage, "sqlite_path", db_path)
    return att, uploads, assets, db_path


def _docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("A paragraph about multimodal retrieval.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "0.95"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_upload_saves_original_and_sidecar_without_vision(monkeypatch, tmp_path):
    att, uploads, _assets, _db = _patch_paths(monkeypatch, tmp_path)

    # Upload must not import/call the VLM path at all.
    import core.multimodal.analyzers as analyzers
    monkeypatch.setattr(
        analyzers, "get_vision_client",
        lambda: (_ for _ in ()).throw(AssertionError("VLM called during upload")),
    )

    rec = att.save_attachment(PNG, "plot.png", content_type="image/png")
    assert (uploads / f"{rec['id']}.png").read_bytes() == PNG
    assert (uploads / f"{rec['id']}.txt").read_text(encoding="utf-8") == ""
    assert rec["multimodal_status"] == "pending"
    assert rec["preview_url"].endswith(f"/{rec['id']}/raw")


def test_supported_image_formats_are_pending(monkeypatch, tmp_path):
    att, _uploads, _assets, _db = _patch_paths(monkeypatch, tmp_path)
    samples = {
        "png": PNG,
        "jpg": b"\xff\xd8\xff\xe0jpeg",
        "webp": b"RIFF\x04\x00\x00\x00WEBPdata",
    }
    for ext, raw in samples.items():
        rec = att.save_attachment(raw, f"figure.{ext}")
        assert rec["ext"] == ext
        assert rec["multimodal_status"] == "pending"
        assert rec["char_count"] == 0


def test_docx_extracts_paragraphs_and_tables(monkeypatch, tmp_path):
    att, uploads, _assets, _db = _patch_paths(monkeypatch, tmp_path)
    raw = _docx_bytes()
    rec = att.save_attachment(raw, "draft.docx")
    text = (uploads / f"{rec['id']}.txt").read_text(encoding="utf-8")
    assert "multimodal retrieval" in text
    assert "| Metric | Value |" in text
    assert "| Accuracy | 0.95 |" in text
    assert rec["multimodal_status"] == "pending"


def test_docx_media_extracts_images_and_skips_broken_entries(monkeypatch, tmp_path):
    att, _uploads, assets, _db = _patch_paths(monkeypatch, tmp_path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("word/media/image1.png", PNG)
        zf.writestr("word/media/not-an-image.bin", b"broken")
    elements = att._docx_elements(buf.getvalue(), "upload:abc", str(assets))
    assert len(elements) == 1
    assert elements[0].kind == "figure"
    assert Path(elements[0].asset_path).read_bytes() == PNG


def test_image_understanding_is_lazy_and_session_cached(monkeypatch, tmp_path):
    att, uploads, _assets, db_path = _patch_paths(monkeypatch, tmp_path)
    rec = att.save_attachment(PNG, "architecture.png")
    session = ChatSession(attachments=[rec])

    import agents.reader_agent as ra
    calls = []

    async def fake_understand(paper, parsed, db):
        calls.append(paper.id)
        parsed.elements[0].understanding = {"description": "encoder and decoder"}
        db.save_elements(paper.id, parsed.elements, parsed.doc_fingerprint)

    monkeypatch.setattr(ra, "_understand_and_persist_elements", fake_understand)

    class FakeStore:
        def upsert_text_chunks(self, *args, **kwargs):
            return 1

    import tools.storage.vectorstore as vsm
    monkeypatch.setattr(vsm, "VectorStore", FakeStore)

    first = asyncio.run(att.ensure_attachment_understood(rec, session, focus="解释结构"))
    second = asyncio.run(att.ensure_attachment_understood(rec, session, focus="再解释"))

    assert first["status"] == second["status"] == "ready"
    assert calls == [f"upload:{rec['id']}"]
    assert rec["element_count"] == 1
    assert "encoder and decoder" in (uploads / f"{rec['id']}.txt").read_text("utf-8")
    from tools.storage.database import Database
    db = Database(db_path)
    try:
        assert len(db.get_elements(f"upload:{rec['id']}")) == 1
    finally:
        db.close()


def test_pdf_understanding_reuses_shared_reader_chokepoint(monkeypatch, tmp_path):
    att, uploads, _assets, _db = _patch_paths(monkeypatch, tmp_path)
    rec = att.save_attachment(b"%PDF-1.4\nnot a complete pdf", "scan.pdf")
    session = ChatSession(attachments=[rec])

    import agents.reader_agent as ra
    calls = []

    async def fake_parse(paper, db, *, assets_dir=None):
        calls.append((paper.id, paper.pdf_path, assets_dir))
        return ParsedPaperDocument(
            raw_text="Recovered scanned page text.",
            elements=[PaperElement(
                element_id=f"{paper.id}::figure::1", kind="figure", ordinal=1,
                caption="Recovered diagram", understanding={"description": "flow"},
            )],
            doc_fingerprint="fp",
        )

    monkeypatch.setattr(ra, "parse_and_understand", fake_parse)

    class FakeStore:
        def upsert_text_chunks(self, *args, **kwargs):
            return 1

    import tools.storage.vectorstore as vsm
    monkeypatch.setattr(vsm, "VectorStore", FakeStore)

    result = asyncio.run(att.ensure_attachment_understood(rec, session))
    again = asyncio.run(att.ensure_attachment_understood(rec, session))
    assert result["status"] == again["status"] == "ready"
    assert len(calls) == 1
    assert calls[0][0] == f"upload:{rec['id']}"
    assert "Recovered scanned page text" in (uploads / f"{rec['id']}.txt").read_text("utf-8")


def test_text_attachment_never_enters_vision_path(monkeypatch, tmp_path):
    att, _uploads, _assets, _db = _patch_paths(monkeypatch, tmp_path)
    rec = att.save_attachment(b"plain notes", "notes.md")
    session = ChatSession(attachments=[rec])

    import agents.reader_agent as ra
    monkeypatch.setattr(
        ra, "parse_and_understand",
        lambda *a, **k: (_ for _ in ()).throw(AssertionError("text file entered VLM path")),
    )
    result = asyncio.run(att.ensure_attachment_understood(rec, session))
    assert result == {"id": rec["id"], "status": "text_only", "element_count": 0}


def test_legacy_text_only_pdf_remains_queryable(monkeypatch, tmp_path):
    att, uploads, _assets, _db = _patch_paths(monkeypatch, tmp_path)
    aid = "a" * 32
    uploads.mkdir(parents=True)
    (uploads / f"{aid}.txt").write_text("legacy extracted PDF text", encoding="utf-8")
    rec = {"id": aid, "filename": "legacy.pdf", "char_count": 25}
    result = asyncio.run(att.ensure_attachment_understood(rec, ChatSession(attachments=[rec])))
    assert result["status"] == "legacy_text_only"
    assert (uploads / f"{aid}.txt").read_text("utf-8") == "legacy extracted PDF text"


def test_attachment_by_id_accepts_exact_unique_filename_only():
    from tools.ingest.attachments import attachment_by_id

    first = {"id": "a" * 32, "filename": "Chart.PNG"}
    second = {"id": "b" * 32, "filename": "notes.txt"}
    session = ChatSession(attachments=[first, second])
    assert attachment_by_id(session, first["id"]) is first
    assert attachment_by_id(session, "chart.png") is first
    assert attachment_by_id(session, "missing.png") is None

    duplicate = {"id": "c" * 32, "filename": "chart.png"}
    session.attachments.append(duplicate)
    assert attachment_by_id(session, "chart.png") is None
